import pandas as pd
import numpy as np

# Import openpyxl styling engines
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

input_file = "attendance.dat"
output_file = "cleaned_attendance_by_employee.xlsx"

print("⏳ Offline Engine: Processing log data with strict 1-hour debouncing...")

try:
    # 1. Read the raw text file columns
    # Extracts: Column 0 (ID), Column 1 (Date), Column 2 (Time)
    df = pd.read_csv(
        input_file,
        sep=r'\s+',
        header=None,
        usecols=[0, 1, 2],
        names=["employee_id", "date", "time"],
        dtype=str,
        engine='python'
    )
except Exception as e:
    print(f"❌ Error reading file: {e}")
    exit()

# 2. Clean strings and sort chronologically
df["timestamp"] = pd.to_datetime(df["date"] + " " + df["time"])
df["employee_id"] = df["employee_id"].str.strip()
df = df.sort_values(by=["employee_id", "timestamp"]).reset_index(drop=True)

# --- CORRECTED 1-HOUR DEBOUNCING FILTER ---
cleaned_logs = []
for emp_id, group in df.groupby("employee_id"):
    group = group.sort_values(by="timestamp").reset_index(drop=True)
    
    if len(group) == 0:
        continue
        
    # Always keep the first punch
    last_saved_row = group.iloc[0]
    cleaned_logs.append(last_saved_row)
    
    # Loop through the rest of the punches
    for idx in range(1, len(group)):
        current_row = group.iloc[idx]
        time_gap = (current_row["timestamp"] - last_saved_row["timestamp"]).total_seconds()
        
        if time_gap >= 3600:  # Must be 1 full hour or more apart to be a distinct action
            cleaned_logs.append(current_row)
            last_saved_row = current_row  # Update baseline to this newly saved row
        else:
            # Drop the duplicate / entry error row
            continue

if not cleaned_logs:
    print("❌ Error: No valid logs remaining after data filtering.")
    exit()

df = pd.DataFrame(cleaned_logs).reset_index(drop=True)
final_rows = []

# 3. Run chronological shift-pairing algorithm with sequence protection
for emp_id, group in df.groupby("employee_id"):
    group = group.sort_values(by="timestamp").reset_index(drop=True)
    
    i = 0
    while i < len(group):
        current_clock = group.loc[i, "timestamp"]
        
        # Last log remaining is marked incomplete
        if i == len(group) - 1:
            final_rows.append({
                "employee_id": emp_id,
                "resumption_date": current_clock.date(),
                "resumption_time": current_clock.strftime("%H:%M:%S"),
                "closing_date": "MISSING",
                "closing_time": "MISSING",
                "hours_worked": 0.0,
                "overtime_hours": 0.0,
                "status": "MISSING CLOSING CLOCK"
            })
            break
            
        next_clock = group.loc[i + 1, "timestamp"]
        time_diff_hours = (next_clock - current_clock).total_seconds() / 3600
        
        # RULE 1: If next punch is within 16 hours, pair them as an active shift!
        if time_diff_hours <= 16.0:
            hours_worked = round(time_diff_hours, 2)
            overtime = round(max(0.0, hours_worked - 8.0), 2)
            
            final_rows.append({
                "employee_id": emp_id,
                "resumption_date": current_clock.date(),
                "resumption_time": current_clock.strftime("%H:%M:%S"),
                "closing_date": next_clock.date(),
                "closing_time": next_clock.strftime("%H:%M:%S"),
                "hours_worked": hours_worked,
                "overtime_hours": overtime,
                "status": "OK"
            })
            i += 2  # Skip forward past both paired logs safely
        else:
            # RULE 2: If next punch is > 16 hours away, they forgot to clock out
            final_rows.append({
                "employee_id": emp_id,
                "resumption_date": current_clock.date(),
                "resumption_time": current_clock.strftime("%H:%M:%S"),
                "closing_date": "MISSING",
                "closing_time": "MISSING",
                "hours_worked": 0.0,
                "overtime_hours": 0.0,
                "status": "MISSING CLOSING CLOCK"
            })
            i += 1  # Move to the next log to look for a clean start point

cleaned_df = pd.DataFrame(final_rows)

# 4. Compile Dashboard Payroll Summary Tab
summary_df = cleaned_df.groupby("employee_id").agg(
    total_days_worked=("resumption_date", "count"),
    total_hours_worked=("hours_worked", "sum"),
    total_overtime_hours=("overtime_hours", "sum"),
    incomplete_clocks=("status", lambda x: (x == "MISSING CLOSING CLOCK").sum())
).reset_index()

summary_df["total_hours_worked"] = summary_df["total_hours_worked"].round(2)
summary_df["total_overtime_hours"] = summary_df["total_overtime_hours"].round(2)

# 5. Write outputs and apply Excel formatting
try:
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Payroll Summary", index=False)
        for emp_id, group in cleaned_df.groupby("employee_id"):
            sheet_name = f"Emp_{emp_id}"
            group.to_excel(writer, sheet_name=sheet_name, index=False)
            
        workbook = writer.book
        
        # Design Theme Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Navy Blue
        zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Zebra light gray
        alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft pastel red
        alert_font = Font(name="Calibri", size=11, color="C00000", bold=True)
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for worksheet in workbook.worksheets:
            # FIXED: Explicitly format only the cells in Row 1 as headers
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Format Data Rows
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row), start=2):
                is_even = (row_idx % 2 == 0)
                for cell in row:
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")
                    
                    if is_even:
                        cell.fill = zebra_fill
                        
                    # Highlight missing items / entry errors in pastel red
                    if cell.value in ["MISSING", "MISSING CLOSING CLOCK"] or (worksheet.title == "Payroll Summary" and cell.column == 5 and isinstance(cell.value, (int, float)) and cell.value > 0):
                        cell.fill = alert_fill
                        cell.font = alert_font
            
            # Auto-fit column width padding via cell sequences (tuple error proofed)
            for col_idx, col_cells in enumerate(worksheet.columns, start=1):
                max_len = max(len(str(cell.value or '')) for cell in col_cells)
                col_letter = get_column_letter(col_idx)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    print(f"✅ Success! Local workbook generated cleanly at '{output_file}'.")

except PermissionError:
    print(f"❌ Permission Denied: Close '{output_file}' if it is open in Excel and try again.")

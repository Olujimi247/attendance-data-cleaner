from fastapi import APIRouter, UploadFile, File, HTTPException, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import pandas as pd
import io
import datetime

# Database layer configurations
from app.database import get_db 
from app.models import AttendanceRecord

# openpyxl styling engines
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/attendance", tags=["Attendance Processing"])

@router.post("/upload-dat")
async def upload_attendance_file(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.endswith('.dat'):
        raise HTTPException(status_code=400, detail="Invalid file type. Please upload a .dat file.")
    
    try:
        contents = await file.read()
        data_stream = io.BytesIO(contents)
        df = pd.read_csv(
            data_stream, sep=r'\s+', header=None, usecols=[0, 1, 2],
            names=["employee_id", "date", "time"], dtype=str, engine='python'
        )
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error parsing .dat file: {str(e)}")

    df["timestamp"] = pd.to_datetime(df["date"] + " " + df["time"])
    df["employee_id"] = df["employee_id"].str.strip()
    df = df.sort_values(by=["employee_id", "timestamp"]).reset_index(drop=True)

    cleaned_logs = []
    for emp_id, group in df.groupby("employee_id"):
        group = group.sort_values(by="timestamp").reset_index(drop=True)
        if len(group) == 0:
            continue
        last_saved_row = group.iloc[0]
        cleaned_logs.append(last_saved_row)
        for idx in range(1, len(group)):
            current_row = group.iloc[idx]
            time_gap = (current_row["timestamp"] - last_saved_row["timestamp"]).total_seconds()
            if time_gap >= 3600:
                cleaned_logs.append(current_row)
                last_saved_row = current_row

    if not cleaned_logs:
        raise HTTPException(status_code=400, detail="No valid logs found after debouncing.")

    df = pd.DataFrame(cleaned_logs).reset_index(drop=True)
    final_rows = []

    for emp_id, group in df.groupby("employee_id"):
        group = group.sort_values(by="timestamp").reset_index(drop=True)
        i = 0
        while i < len(group):
            current_clock = group.loc[i, "timestamp"]
            if i == len(group) - 1:
                final_rows.append({
                    "employee_id": emp_id, "resumption_date": current_clock.date(),
                    "resumption_time": current_clock.strftime("%H:%M:%S"), "closing_date": "MISSING",
                    "closing_time": "MISSING", "hours_worked": 0.0, "overtime_hours": 0.0,
                    "status": "MISSING CLOSING CLOCK"
                })
                break
            next_clock = group.loc[i + 1, "timestamp"]
            time_diff_hours = (next_clock - current_clock).total_seconds() / 3600
            if time_diff_hours <= 16.0:
                hours_worked = round(time_diff_hours, 2)
                overtime = round(max(0.0, hours_worked - 8.0), 2)
                final_rows.append({
                    "employee_id": emp_id, "resumption_date": current_clock.date(),
                    "resumption_time": current_clock.strftime("%H:%M:%S"), "closing_date": next_clock.date(),
                    "closing_time": next_clock.strftime("%H:%M:%S"), "hours_worked": hours_worked,
                    "overtime_hours": overtime, "status": "OK"
                })
                i += 2
            else:
                final_rows.append({
                    "employee_id": emp_id, "resumption_date": current_clock.date(),
                    "resumption_time": current_clock.strftime("%H:%M:%S"), "closing_date": "MISSING",
                    "closing_time": "MISSING", "hours_worked": 0.0, "overtime_hours": 0.0,
                    "status": "MISSING CLOSING CLOCK"
                })
                i += 1

    cleaned_df = pd.DataFrame(final_rows)

    try:
        for _, row in cleaned_df.iterrows():
            existing_record = db.query(AttendanceRecord).filter(
                AttendanceRecord.employee_id == str(row["employee_id"]),
                AttendanceRecord.resumption_date == row["resumption_date"]
            ).first()
            res_time_obj = datetime.time.fromisoformat(row["resumption_time"])
            if existing_record:
                existing_record.resumption_time = res_time_obj
                existing_record.closing_date = str(row["closing_date"])
                existing_record.closing_time = str(row["closing_time"])
                existing_record.hours_worked = float(row["hours_worked"])
                existing_record.overtime_hours = float(row["overtime_hours"])
                existing_record.status = str(row["status"])
            else:
                new_record = AttendanceRecord(
                    employee_id=str(row["employee_id"]), resumption_date=row["resumption_date"],
                    resumption_time=res_time_obj, closing_date=str(row["closing_date"]),
                    closing_time=str(row["closing_time"]), hours_worked=float(row["hours_worked"]),
                    overtime_hours=float(row["overtime_hours"]), status=str(row["status"])
                )
                db.add(new_record)
        db.commit()
        return {
            "status": "success",
            "message": f"Successfully parsed raw log data and synced {len(cleaned_df)} items."
        }
    except Exception as db_err:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Database synchronization error: {str(db_err)}")

# --- CRITICAL FIX: PLACED OUTSIDE THE POST ENDPOINT FUNCTION BLOCK ---
@router.get("/download-report")
def download_payroll_report(db: Session = Depends(get_db)):
    records = db.query(AttendanceRecord).all()
    if not records:
        raise HTTPException(status_code=404, detail="No attendance logs found in the database.")
        
    data = [{
        "employee_id": r.employee_id,
        "resumption_date": r.resumption_date,
        "resumption_time": r.resumption_time.strftime("%H:%M:%S") if r.resumption_time else "MISSING",
        "closing_date": r.closing_date,
        "closing_time": r.closing_time,
        "hours_worked": r.hours_worked,
        "overtime_hours": r.overtime_hours,
        "status": r.status
    } for r in records]
    
    cleaned_df = pd.DataFrame(data)
    summary_df = cleaned_df.groupby("employee_id").agg(
        total_days_worked=("resumption_date", "count"),
        total_hours_worked=("hours_worked", "sum"),
        total_overtime_hours=("overtime_hours", "sum"),
        incomplete_clocks=("status", lambda x: (x == "MISSING CLOSING CLOCK").sum())
    ).reset_index()
    
    summary_df["total_hours_worked"] = summary_df["total_hours_worked"].round(2)
    summary_df["total_overtime_hours"] = summary_df["total_overtime_hours"].round(2)
    
    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Payroll Summary", index=False)
        for emp_id, group in cleaned_df.groupby("employee_id"):
            sheet_name = f"Emp_{emp_id}"
            group.to_excel(writer, sheet_name=sheet_name, index=False)
            
        workbook = writer.book
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        alert_font = Font(name="Calibri", size=11, color="C00000", bold=True)
        thin_border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
        
        for worksheet in workbook.worksheets:
            for col in worksheet.iter_cols(min_row=1, max_row=1):
                for cell in col:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row), start=2):
                is_even = (row_idx % 2 == 0)
                for cell in row:
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")
                    if is_even:
                        cell.fill = zebra_fill
                    if cell.value in ["MISSING", "MISSING CLOSING CLOCK"] or (worksheet.title == "Payroll Summary" and cell.column == 5 and isinstance(cell.value, (int, float)) and cell.value > 0):
                        cell.fill = alert_fill
                        cell.font = alert_font
                        
            for col_idx, col_cells in enumerate(worksheet.columns, start=1):
                max_len = max(len(str(cell.value or '')) for cell in col_cells)
                col_letter = get_column_letter(col_idx)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
                
    excel_buffer.seek(0)
    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=final_payroll_report.xlsx"}
    )

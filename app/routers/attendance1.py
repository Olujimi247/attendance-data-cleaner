from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
import pandas as pd
import io

# Import openpyxl styling engines
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/attendance", tags=["Attendance Processing"])

@router.post("/upload-dat")
async def upload_attendance_file(file: UploadFile = File(...)):
    # 1. Ensure file type constraint is satisfied
    if not file.filename.endswith('.dat'):
        raise HTTPException(status_code=400, detail="Invalid file type. Please upload a .dat file.")
    
    try:
        # 2. Read the binary raw stream directly into Pandas
        contents = await file.read()
        data_stream = io.BytesIO(contents)
        
        # Extracts: Column 0 (ID), Column 1 (Date), Column 2 (Time)
        df = pd.read_csv(
            data_stream,
            sep=r'\s+',
            header=None,
            usecols=[0, 1, 2],
            names=["employee_id", "date", "time"],
            dtype=str,
            engine='python'
        )
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error parsing .dat file structure: {str(e)}")

    # 3. Clean strings and sort chronologically
    df["timestamp"] = pd.to_datetime(df["date"] + " " + df["time"])
    df["employee_id"] = df["employee_id"].str.strip()
    df = df.sort_values(by=["employee_id", "timestamp"]).reset_index(drop=True)

    # --- 1-HOUR ACCIDENTAL DOUBLE-PRESS (DEBOUNCING) FILTER ---
    # Drops any punch that happens within 60 minutes (3600 seconds) of the last recorded punch
    cleaned_logs = []
    for emp_id, group in df.groupby("employee_id"):
        group = group.sort_values(by="timestamp").reset_index(drop=True)
        last_saved_time = None
        
        for idx, row in group.iterrows():
            if last_saved_time is None:
                cleaned_logs.append(row)
                last_saved_time = row["timestamp"]
            else:
                time_gap = (row["timestamp"] - last_saved_time).total_seconds()
                if time_gap >= 3600:  # Must be 1 full hour or more to count as a separate action
                    cleaned_logs.append(row)
                    last_saved_time = row["timestamp"]
                else:
                    # Ignore the duplicate double press
                    continue

    if not cleaned_logs:
        raise HTTPException(status_code=400, detail="No valid logs found after data cleanup filtering.")

    df = pd.DataFrame(cleaned_logs).reset_index(drop=True)
    final_rows = []

    # 4. Run chronological shift-pairing algorithm with sequence protection
    for emp_id, group in df.groupby("employee_id"):
        group = group.sort_values(by="timestamp").reset_index(drop=True)
        
        i = 0
        while i < len(group):
            current_clock = group.loc[i, "timestamp"]
            
            # Last log remaining is marked incomplete
            if i == len(group) - 1:
                final_rows.append({
                    "employee_id": emp_id,
                    "resumption_date": current_clock.date(),
                    "resumption_time": current_clock.strftime("%H:%M:%S"),
                    "closing_date": "MISSING",
                    "closing_time": "MISSING",
                    "hours_worked": 0.0,
                    "overtime_hours": 0.0,
                    "status": "MISSING CLOSING CLOCK"
                })
                break
                
            next_clock = group.loc[i + 1, "timestamp"]
            time_diff_hours = (next_clock - current_clock).total_seconds() / 3600
            
            # RULE 1: If next punch is within 16 hours, pair them as an active shift!
            if time_diff_hours <= 16.0:
                hours_worked = round(time_diff_hours, 2)
                overtime = round(max(0.0, hours_worked - 8.0), 2)
                
                final_rows.append({
                    "employee_id": emp_id,
                    "resumption_date": current_clock.date(),
                    "resumption_time": current_clock.strftime("%H:%M:%S"),
                    "closing_date": next_clock.date(),
                    "closing_time": next_clock.strftime("%H:%M:%S"),
                    "hours_worked": hours_worked,
                    "overtime_hours": overtime,
                    "status": "OK"
                })
                i += 2  # Skip forward past both paired logs safely
            else:
                # RULE 2: If next punch is > 16 hours away, they forgot to clock out
                final_rows.append({
                    "employee_id": emp_id,
                    "resumption_date": current_clock.date(),
                    "resumption_time": current_clock.strftime("%H:%M:%S"),
                    "closing_date": "MISSING",
                    "closing_time": "MISSING",
                    "hours_worked": 0.0,
                    "overtime_hours": 0.0,
                    "status": "MISSING CLOSING CLOCK"
                })
                i += 1  # Move to the next log to look for a clean start point

    cleaned_df = pd.DataFrame(final_rows)

    # 5. Compile Dashboard Payroll Summary Tab
    summary_df = cleaned_df.groupby("employee_id").agg(
        total_days_worked=("resumption_date", "count"),
        total_hours_worked=("hours_worked", "sum"),
        total_overtime_hours=("overtime_hours", "sum"),
        incomplete_clocks=("status", lambda x: (x == "MISSING CLOSING CLOCK").sum())
    ).reset_index()

    summary_df["total_hours_worked"] = summary_df["total_hours_worked"].round(2)
    summary_df["total_overtime_hours"] = summary_df["total_overtime_hours"].round(2)

    # 6. Stream outputs into memory and apply Excel formatting
    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Payroll Summary", index=False)
        for emp_id, group in cleaned_df.groupby("employee_id"):
            sheet_name = f"Emp_{emp_id}"
            group.to_excel(writer, sheet_name=sheet_name, index=False)
            
        workbook = writer.book
        
        # Design Theme Objects
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Navy Blue
        zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Zebra light gray
        alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Light pastel red
        alert_font = Font(name="Calibri", size=11, color="C00000", bold=True)
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for worksheet in workbook.worksheets:
            # Format Column Row Headers
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Format Data Elements Rows
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row), start=2):
                is_even = (row_idx % 2 == 0)
                for cell in row:
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")
                    
                    if is_even:
                        cell.fill = zebra_fill
                        
                    # Highlight human errors / missing entries in soft red
                    if cell.value in ["MISSING", "MISSING CLOSING CLOCK"] or (worksheet.title == "Payroll Summary" and cell.column == 5 and isinstance(cell.value, (int, float)) and cell.value > 0):
                        cell.fill = alert_fill
                        cell.font = alert_font
            
            # FIXED: Auto-fit column widths using enumerated cell sequences to avoid tuple errors
            for col_idx, col_cells in enumerate(worksheet.columns, start=1):
                max_len = max(len(str(cell.value or '')) for cell in col_cells)
                col_letter = get_column_letter(col_idx)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel_buffer.seek(0)
    
    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=final_payroll_report.xlsx"}
    )
